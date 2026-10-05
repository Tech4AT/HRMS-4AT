"""Fills an HR-uploaded .docx offer-letter template and converts it to PDF.

Uses `docxtpl` (Jinja2-in-Word) to substitute `{{placeholder}}` tokens typed
directly into the document — it handles Word's habit of splitting one piece
of visible text across several XML runs, which naive find/replace on the
raw XML does not. The filled .docx is then converted to PDF using LibreOffice's
headless `--convert-to pdf` mode, which works on Ubuntu/Linux servers.
"""
import logging
import tempfile
from pathlib import Path

from docxtpl import DocxTemplate

logger = logging.getLogger(__name__)


class DocxRenderError(ValueError):
    """Raised when the uploaded template can't be filled or converted —
    e.g. it isn't a valid .docx, or LibreOffice isn't available on this
    machine. Subclasses ValueError so it's caught by the same handling as
    other offer-letter validation errors (see serializers.py, views.py)
    rather than surfacing as a raw 500."""


def _convert_via_libreoffice(docx_path: str, pdf_path: str) -> None:
    """Convert a .docx to PDF using LibreOffice headless mode.

    LibreOffice writes the output PDF to the same directory as the input
    file, naming it after the input stem (rendered.docx -> rendered.pdf),
    which matches the pdf_path the caller expects."""
    import subprocess
    from pathlib import Path

    out_dir = str(Path(docx_path).parent)
    result = subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', out_dir, docx_path],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr or result.stdout or 'libreoffice conversion failed')


def _expand_salary_table(docx_path: Path, rows: list) -> None:
    """Templates written before the payroll engine have a fixed salary table
    (Basic Salary / HRA / Other Allowances / Other Components / Total CTC).
    Swap its component rows for the real breakup rows so the letter shows what
    payroll will actually pay. The template file itself is never modified —
    this runs on the temporary copy. Tables that don't match are left alone."""
    import copy

    from docx import Document
    from docx.table import _Row

    document = Document(str(docx_path))
    changed = False
    for table in document.tables:
        labels = [row.cells[0].text.strip() for row in table.rows]
        if 'Basic Salary' not in labels or 'Total CTC' not in labels:
            continue
        start, end = labels.index('Basic Salary'), labels.index('Total CTC')
        if end <= start or len(table.rows[start].cells) < 3:
            continue
        proto = table.rows[start]._tr
        anchor = table.rows[end]._tr
        old = [row._tr for row in table.rows[start:end]]
        for name, monthly, annual in rows:
            tr = copy.deepcopy(proto)
            anchor.addprevious(tr)
            for cell, text in zip(_Row(tr, table).cells, (name, monthly, annual)):
                paragraphs = cell.paragraphs
                runs = paragraphs[0].runs
                if runs:
                    runs[0].text = text
                    for extra in runs[1:]:
                        extra.text = ''
                else:
                    paragraphs[0].add_run(text)
                for extra in paragraphs[1:]:
                    extra._p.getparent().remove(extra._p)
        for tr in old:
            tr.getparent().remove(tr)
        changed = True
    if changed:
        document.save(str(docx_path))


def _append_signature_block(docx_path: Path, context: dict) -> None:
    """Mirrors `_render_text_template_pdf`'s post-signature block for the
    reportlab path (offer_letter.py) — an uploaded .docx template has no
    `{{signature_name}}`/`{{signed_date}}` tokens of its own (HR never typed
    them in), so docxtpl's substitution silently has nothing to fill. This
    appends the same "Signed and Accepted" section directly, so a signed
    offer looks signed regardless of which renderer produced it."""
    from docx import Document

    document = Document(str(docx_path))
    document.add_paragraph()
    document.add_heading('Signed and Accepted', level=2)
    document.add_paragraph(f"Signature: {context.get('signature_name', '')}")
    document.add_paragraph(f"Date: {context.get('signed_date', '')}")
    document.save(str(docx_path))


def render_docx_template_to_pdf(source_path: str, context: dict, *, signed: bool = False) -> bytes:
    with tempfile.TemporaryDirectory(prefix='offer_letter_') as tmp_dir:
        tmp_dir_path = Path(tmp_dir)
        try:
            source = source_path
            rows = context.get('salary_rows') or []
            if rows:
                prepared = tmp_dir_path / 'prepared.docx'
                import shutil

                shutil.copyfile(source_path, prepared)
                _expand_salary_table(prepared, rows)
                source = str(prepared)
            tpl = DocxTemplate(source)
            tpl.render(context)
        except Exception as exc:  # docxtpl/Jinja2 raise several distinct types
            raise DocxRenderError(f'Could not fill the Word template: {exc}') from exc

        rendered_docx = tmp_dir_path / 'rendered.docx'
        tpl.save(str(rendered_docx))

        if signed:
            _append_signature_block(rendered_docx, context)

        rendered_pdf = tmp_dir_path / 'rendered.pdf'
        try:
            _convert_via_libreoffice(str(rendered_docx), str(rendered_pdf))
        except Exception as exc:
            logger.exception('libreoffice conversion failed')
            raise DocxRenderError(
                'Could not convert the Word template to PDF. This requires LibreOffice '
                'to be installed on the server — verify it is available with `libreoffice --version`.'
            ) from exc

        if not rendered_pdf.exists():
            raise DocxRenderError('LibreOffice did not produce a PDF file — conversion failed silently.')

        return rendered_pdf.read_bytes()
